# -*- coding: utf-8 -*-
"""资料助手核心逻辑：解析文件 -> SQLite + FTS5 索引 -> 搜索 -> 表格匹配填写。

只依赖 Python 标准库 + 可选解析库（pandas/openpyxl/python-docx/python-pptx/pdfplumber/PIL）。
库缺失时自动降级（txt/csv 仍可用，其他格式会提示错误）。
"""
from __future__ import annotations

import csv
import datetime as _dt
import difflib
import hashlib
import html
import io
import json
import os
import re
import sqlite3
import time
import unicodedata
from typing import Any, Dict, List, Optional

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    import pandas as pd
except Exception:
    pd = None

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    import docx as _docx_mod
except Exception:
    _docx_mod = None

try:
    import pptx as _pptx_mod
except Exception:
    _pptx_mod = None

try:
    from PIL import Image
except Exception:
    Image = None

try:
    import pytesseract
except Exception:
    pytesseract = None

try:
    import xlrd  # 旧版 .xls (BIFF)
except Exception:
    xlrd = None

try:
    import olefile  # 旧版 .doc/.ppt (OLE 复合文档)
except Exception:
    olefile = None

DB_PATH = "data/kb.sqlite3"
FILES_DIR = "data/files"
EXPORT_DIR = "data/export"

TEXT_EXTS = {".txt", ".md", ".markdown", ".log", ".json", ".xml", ".yaml", ".yml", ".ini", ".conf"}
TABLE_EXTS = {".csv", ".xlsx", ".xlsm", ".xls"}
DOC_EXTS = {".docx", ".doc", ".rtf"}
PDF_EXTS = {".pdf"}
PPT_EXTS = {".pptx", ".ppt"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".jfif"}
HTML_EXTS = {".html", ".htm"}


def now_str() -> str:
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _safe_rel_dir(folder: Any) -> str:
    """把前端传来的相对目录清洗成安全路径（去掉 .. / . / 空段 / 斜杠差异）。"""
    if not folder:
        return ""
    parts = []
    for p in str(folder).replace("\\", "/").split("/"):
        p = p.strip()
        if p in ("", ".", ".."):
            continue
        parts.append(p)
    return "/".join(parts)


def norm_text(s: Any) -> str:
    """规范化：NFKC 统一字符、去空白，用于精确比较。"""
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKC", s)
    s = s.lower()
    s = re.sub(r"[\s\u3000]+", "", s)
    return s


CJK_RE = re.compile(r"[\u4e00-\u9fff\u3400-\u4dbf]+")


def cjk_tokens(text: str) -> List[str]:
    """中文按 2 字 bigram 切分，英文/数字按单词切分，用于 FTS 索引与查询。"""
    out: List[str] = []
    text = unicodedata.normalize("NFKC", str(text))
    for tok in re.findall(r"[A-Za-z0-9][A-Za-z0-9_.\-]*", text):
        out.append(tok.lower())
    for run in CJK_RE.findall(text):
        if len(run) == 1:
            out.append(run)
        else:
            for i in range(len(run) - 1):
                out.append(run[i:i + 2])
    return out


def fts_index_text(text: str) -> str:
    return " ".join(cjk_tokens(text))


def highlight(text: str, query: str, max_len: int = 220) -> str:
    """在片段中高亮命中的关键字，返回 HTML。"""
    if not text:
        return ""
    text = html.escape(text)
    q_norm = norm_text(query)
    t_norm = norm_text(text)
    pos = t_norm.find(q_norm) if q_norm else -1
    if pos < 0:
        for ch in query[:8]:
            p = t_norm.find(norm_text(ch))
            if p >= 0:
                pos = p
                break
    if pos < 0:
        pos = 0
    start = max(0, pos - max_len // 3)
    seg = text[start:start + max_len]
    if start > 0:
        seg = "\u2026" + seg
    if start + max_len < len(text):
        seg = seg + "\u2026"
    for term in re.split(r"[\s,，。;；、]+", query.strip())[:8]:
        if not term:
            continue
        pat = re.compile(re.escape(html.escape(term)), re.IGNORECASE)
        seg = pat.sub(lambda m: "<mark>" + m.group(0) + "</mark>", seg)
    return seg


def fmt_size(n: int) -> str:
    if n is None:
        return ""
    n = float(n)
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} GB"


class ParseResult:
    """一个文件解析出的多条记录。"""

    def __init__(self, title: str, content: str, rtype: str = "doc",
                 page: Optional[int] = None, row: Optional[int] = None,
                 fields: Optional[Dict[str, Any]] = None, extra: Optional[Dict[str, Any]] = None):
        self.title = (title or "").strip()
        self.content = content or ""
        self.rtype = rtype
        self.page = page
        self.row = row
        self.fields = fields or {}
        self.extra = extra or {}


def _clean_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _dedup_headers(headers: List[str]) -> List[str]:
    seen: Dict[str, int] = {}
    out = []
    for h in headers:
        h = _clean_cell(h) or ("列%d" % (len(out) + 1))
        if h in seen:
            seen[h] += 1
            h = "%s_%d" % (h, seen[h])
        else:
            seen[h] = 1
        out.append(h)
    return out


TITLE_KEYWORDS = {
    "名称": 3, "标题": 3, "品名": 3, "姓名": 3, "主题": 3, "subject": 3,
    "name": 3, "title": 3, "产品": 1, "项目": 1, "型号": 1, "品牌": 1, "科目": 1,
}


def _pick_title_col(headers: List[str]) -> int:
    """从表头里挑一个更像“标题/名称”的列作为记录标题来源，默认第一列。"""
    best, best_rank = 0, -1
    for i, h in enumerate(headers):
        nh = norm_text(h)
        rank = sum(w for k, w in TITLE_KEYWORDS.items() if k in nh)
        if rank > best_rank:
            best, best_rank = i, rank
    return best

def _rows_to_records(rows: List[List[str]], fname: str, rtype: str) -> List[ParseResult]:
    """把二维表格转成记录（第一行有内容且列数>=2 的作为表头）。"""
    header_idx = -1
    for i, r in enumerate(rows[:20]):
        cells = [_clean_cell(c) for c in r]
        if sum(1 for c in cells if c) >= 2:
            header_idx = i
            break
    if header_idx < 0:
        body = "\n".join("\t".join(_clean_cell(c) for c in r) for r in rows if any(_clean_cell(c) for c in r))
        return [ParseResult(fname, body, rtype)]
    headers = _dedup_headers([_clean_cell(c) for c in rows[header_idx]])
    title_idx = _pick_title_col(headers)
    results: List[ParseResult] = []
    for ri, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        cells = [_clean_cell(c) for c in r] + [""] * (len(headers) - len(r))
        if not any(cells):
            continue
        fields = {h: cells[i] if i < len(cells) else "" for i, h in enumerate(headers)}
        title = fields.get(headers[title_idx]) or fields.get(headers[0]) or ("%s 第%d行" % (fname, ri))
        content = " ".join("%s %s" % (h, v) for h, v in fields.items() if v)
        results.append(ParseResult(title, content, rtype, row=ri, fields=fields,
                                   extra={"sheet": "", "row": ri, "headers": headers}))
    return results


def parse_text_file(path: str) -> List[ParseResult]:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    return [ParseResult(os.path.basename(path), content, "text")]


def parse_html_file(path: str) -> List[ParseResult]:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read()
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    text = re.sub(r"(?s)<[^>]+>", " ", raw)
    text = html.unescape(re.sub(r"\s+", " ", text)).strip()
    return [ParseResult(os.path.basename(path), text, "text")]


def parse_csv_file(path: str) -> List[ParseResult]:
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return _rows_to_records(rows, os.path.basename(path), "sheet")


def parse_xlsx_file(path: str) -> List[ParseResult]:
    if openpyxl is None:
        return [ParseResult(os.path.basename(path), "", "sheet", extra={"error": "缺少 openpyxl 库"})]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    results: List[ParseResult] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_clean_cell(c) for c in row])
        recs = _rows_to_records(rows, os.path.basename(path), "sheet")
        for rec in recs:
            rec.extra["sheet"] = ws.title
        results.extend(recs)
    wb.close()
    return results
def parse_xls_file(path: str) -> List[ParseResult]:
    """旧版 .xls：用 xlrd 读取（每行 = 一条记录）。"""
    if xlrd is None:
        return [ParseResult(os.path.basename(path), "", "sheet", extra={"error": "缺少 xlrd 库"})]
    try:
        wb = xlrd.open_workbook(path)
    except Exception as e:
        return [ParseResult(os.path.basename(path), "", "sheet", extra={"error": "解析 .xls 失败: %s" % e})]
    results: List[ParseResult] = []
    for sh in wb.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                row.append(_xlrd_cell_value(sh.cell(r, c), wb))
            rows.append(row)
        recs = _rows_to_records(rows, os.path.basename(path), "sheet")
        for rec in recs:
            rec.extra["sheet"] = sh.name
        results.extend(recs)
    return results


def _xlrd_cell_value(cell, wb) -> str:
    t = cell.ctype
    if t == xlrd.XL_CELL_TEXT:
        return str(cell.value).strip()
    if t == xlrd.XL_CELL_NUMBER:
        v = cell.value
        return str(int(v)) if float(v).is_integer() else str(v)
    if t == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, wb.datemode).strftime("%Y-%m-%d")
        except Exception:
            return str(cell.value)
    if t == xlrd.XL_CELL_BOOLEAN:
        return "是" if cell.value else "否"
    return ""


def parse_doc_file(path: str) -> List[ParseResult]:
    """旧版 .doc：尽力从 WordDocument 流提取正文（best-effort）。"""
    if olefile is None:
        return [ParseResult(os.path.basename(path), "", "doc", extra={"error": "缺少 olefile 库"})]
    text = ""
    extra: Dict[str, Any] = {}
    try:
        ole = olefile.OleFileIO(path)
        try:
            if ole.exists("WordDocument"):
                data = ole.openstream("WordDocument").read()
                fc_min = int.from_bytes(data[0x18:0x1C], "little")
                fc_max = int.from_bytes(data[0x1C:0x20], "little")
                if 0 <= fc_min <= fc_max <= len(data):
                    text = _decode_word_text(data[fc_min:fc_max])
                    extra["method"] = "fib"
        finally:
            ole.close()
    except Exception as e:
        extra["error"] = "解析 .doc 失败: %s" % e
    return [ParseResult(os.path.basename(path), text.strip(), "doc", extra=extra)]


def _decode_word_text(chunk: bytes) -> str:
    """Word 97+ 正文一般是 UTF-16LE，乱码时再尝试 GBK/CP1252。"""
    cands = []
    try:
        t = chunk.decode("utf-16-le")
        cands.append((t, _text_quality(t)))
    except Exception:
        pass
    for enc in ("gb18030", "cp1252"):
        try:
            t = chunk.decode(enc)
            cands.append((t, _text_quality(t)))
        except Exception:
            pass
    if not cands:
        return ""
    cands.sort(key=lambda x: -x[1])
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]+", "\n", cands[0][0])


def _text_quality(s: str) -> float:
    if not s:
        return 0.0
    printable = sum(1 for ch in s if ch.isprintable() and not ch.isspace())
    return printable / max(len(s), 1)


def parse_ppt_file(path: str) -> List[ParseResult]:
    """旧版 .ppt：遍历 PowerPoint Document 流里的文本原子（best-effort）。"""
    if olefile is None:
        return [ParseResult(os.path.basename(path), "", "slide", extra={"error": "缺少 olefile 库"})]
    parts: List[str] = []
    extra: Dict[str, Any] = {}
    try:
        ole = olefile.OleFileIO(path)
        try:
            if ole.exists("PowerPoint Document"):
                data = ole.openstream("PowerPoint Document").read()
                _walk_ppt_records(data, 0, len(data), parts)
                extra["method"] = "record"
        finally:
            ole.close()
    except Exception as e:
        extra["error"] = "解析 .ppt 失败: %s" % e
    return [ParseResult(os.path.basename(path), "\n".join(parts).strip(), "slide", extra=extra)]


def _walk_ppt_records(data: bytes, start: int, end: int, out: List[str]):
    off = start
    while off + 8 <= end:
        ver = data[off] & 0x0F
        rtype = int.from_bytes(data[off + 2:off + 4], "little")
        rlen = int.from_bytes(data[off + 4:off + 8], "little")
        payload = data[off + 8:off + 8 + rlen]
        if ver == 0x0F:  # 容器记录：递归
            _walk_ppt_records(data, off + 8, off + 8 + rlen, out)
        else:
            if rtype == 0x0FA0:  # TextCharsAtom：UTF-16LE
                t = payload.decode("utf-16-le", errors="replace").rstrip("\x00")
                if t.strip():
                    out.append(t)
            elif rtype == 0x0FA8:  # TextBytesAtom：系统 ANSI(中文多为 GBK)
                t = payload.decode("gb18030", errors="replace").rstrip("\x00")
                if t.strip():
                    out.append(t)
        off += 8 + rlen

def parse_docx_file(path: str) -> List[ParseResult]:
    if _docx_mod is None:
        return [ParseResult(os.path.basename(path), "", "doc", extra={"error": "缺少 python-docx 库"})]
    d = _docx_mod.Document(path)
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    results: List[ParseResult] = []
    doc_text = "\n".join(paras)
    for ti, table in enumerate(d.tables):
        grid = [[_clean_cell(c.text) for c in row.cells] for row in table.rows]
        if not grid:
            continue
        if max(len(r) for r in grid) <= 2:
            # 两列参数表（如规格表）并入正文
            lines = ["\t".join(r) for r in grid if any(r)]
            doc_text = (doc_text + "\n" + "\n".join(lines)).strip()
            continue
        recs = _rows_to_records(grid, "%s\u00b7表%d" % (os.path.basename(path), ti + 1), "sheet")
        for rec in recs:
            rec.extra["table"] = ti + 1
            rec.extra["sheet"] = "表%d" % (ti + 1)
        results.extend(recs)
    results.insert(0, ParseResult(os.path.basename(path), doc_text, "doc"))
    return results


def parse_pptx_file(path: str) -> List[ParseResult]:
    if _pptx_mod is None:
        return [ParseResult(os.path.basename(path), "", "slide", extra={"error": "缺少 python-pptx 库"})]
    prs = _pptx_mod.Presentation(path)
    results: List[ParseResult] = []
    for si, slide in enumerate(prs.slides, start=1):
        texts = []
        title = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip()).strip()
            if t:
                texts.append(t)
                if not title:
                    title = t.splitlines()[0][:60]
        results.append(ParseResult(title or ("%s 第%d页" % (os.path.basename(path), si)),
                                   "\n".join(texts), "slide", page=si))
    return results


def parse_pdf_file(path: str) -> List[ParseResult]:
    if pdfplumber is None:
        return [ParseResult(os.path.basename(path), "", "pdf", extra={"error": "缺少 pdfplumber 库"})]
    results: List[ParseResult] = []
    with pdfplumber.open(path) as pdf:
        for pi, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            try:
                tables = page.extract_tables() or []
            except Exception:
                tables = []
            for table in tables:
                grid = [[_clean_cell(c) for c in row] for row in table]
                grid = [r for r in grid if any(r)]
                if not grid:
                    continue
                if max(len(r) for r in grid) <= 2:
                    lines = ["\t".join(r) for r in grid]
                    if text:
                        text += "\n"
                    text += "\n".join(lines)
                    continue
                recs = _rows_to_records(grid, "%s 第%d页表格" % (os.path.basename(path), pi), "sheet")
                for rec in recs:
                    rec.page = pi
                results.extend(recs)
            if text:
                results.append(ParseResult("%s 第%d页" % (os.path.basename(path), pi), text, "pdf", page=pi))
    return results


def parse_image_file(path: str) -> List[ParseResult]:
    content = ""
    extra: Dict[str, Any] = {}
    stem = os.path.splitext(path)[0]
    sidecar = stem + ".txt"
    if os.path.exists(sidecar):
        try:
            with open(sidecar, "r", encoding="utf-8", errors="replace") as f:
                content += f.read().strip() + "\n"
        except Exception:
            pass
    if Image is not None:
        try:
            with Image.open(path) as im:
                extra["width"] = im.width
                extra["height"] = im.height
                extra["format"] = im.format or ""
        except Exception:
            pass
    if pytesseract is not None and not content.strip():
        try:
            content = (pytesseract.image_to_string(Image.open(path), lang="chi_sim+eng") or "").strip()
        except Exception:
            pass
    return [ParseResult(os.path.basename(path), content.strip(), "image", fields=extra, extra=extra)]


EXT_PARSERS = {}
for _e in TEXT_EXTS | HTML_EXTS:
    EXT_PARSERS[_e] = parse_html_file if _e in HTML_EXTS else parse_text_file
EXT_PARSERS[".csv"] = parse_csv_file
EXT_PARSERS[".xlsx"] = parse_xlsx_file
EXT_PARSERS[".xlsm"] = parse_xlsx_file
EXT_PARSERS[".xls"] = parse_xls_file
EXT_PARSERS[".docx"] = parse_docx_file
EXT_PARSERS[".doc"] = parse_doc_file
EXT_PARSERS[".pptx"] = parse_pptx_file
EXT_PARSERS[".ppt"] = parse_ppt_file
for _e in PDF_EXTS:
    EXT_PARSERS[_e] = parse_pdf_file


def kind_of(ext: str) -> str:
    ext = ext.lower()
    if ext in TEXT_EXTS or ext in HTML_EXTS:
        return "text"
    if ext in TABLE_EXTS:
        return "sheet"
    if ext in DOC_EXTS:
        return "doc"
    if ext in PDF_EXTS:
        return "pdf"
    if ext in PPT_EXTS:
        return "slide"
    if ext in IMAGE_EXTS:
        return "image"
    return "other"


class KnowledgeBase:
    def __init__(self, root_dir: str):
        self.root = root_dir
        self.files_dir = os.path.join(root_dir, FILES_DIR)
        self.export_dir = os.path.join(root_dir, EXPORT_DIR)
        self.db_path = os.path.join(root_dir, DB_PATH)
        os.makedirs(self.files_dir, exist_ok=True)
        os.makedirs(self.export_dir, exist_ok=True)
        self.conn = sqlite3.connect(self.db_path, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA journal_mode=WAL")
        self._init_db()

    def _init_db(self) -> None:
        c = self.conn
        c.execute("""
            CREATE TABLE IF NOT EXISTS files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                path TEXT UNIQUE, name TEXT, ext TEXT, kind TEXT,
                size INTEGER, mtime REAL, sha TEXT,
                status TEXT DEFAULT 'ok', error TEXT,
                record_count INTEGER DEFAULT 0, created_at TEXT
            )""")
        c.execute("""
            CREATE TABLE IF NOT EXISTS records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                file_id INTEGER, title TEXT, content TEXT, rtype TEXT,
                page INTEGER, row_no INTEGER, fields TEXT, extra TEXT
            )""")
        c.execute("CREATE INDEX IF NOT EXISTS idx_records_file ON records(file_id)")
        c.execute("""
            CREATE VIRTUAL TABLE IF NOT EXISTS fts USING fts5(
                title, content, tokenize='unicode61'
            )""")
        c.execute("PRAGMA foreign_keys=ON")
        self.conn.commit()

    def _insert_file(self, rel_path: str, name: str, ext: str, size: int,
                     mtime: float, sha: str) -> int:
        kind = kind_of(ext)
        cur = self.conn.execute(
            "INSERT OR REPLACE INTO files(path,name,ext,kind,size,mtime,sha,created_at) "
            "VALUES(?,?,?,?,?,?,?,?)",
            (rel_path, name, ext, kind, size, mtime, sha, now_str()))
        self.conn.commit()
        return cur.lastrowid

    def _delete_file(self, file_id: int) -> None:
        with self.conn:
            row = self.conn.execute("SELECT path FROM files WHERE id=?", (file_id,)).fetchone()
            # 先删全文索引与记录，避免残留孤儿数据
            old = self.conn.execute("SELECT id FROM records WHERE file_id=?", (file_id,)).fetchall()
            for o in old:
                self.conn.execute("DELETE FROM fts WHERE rowid=?", (o["id"],))
            self.conn.execute("DELETE FROM records WHERE file_id=?", (file_id,))
            self.conn.execute("DELETE FROM files WHERE id=?", (file_id,))
            if row:
                p = os.path.join(self.root, row["path"])
                try:
                    if os.path.isfile(p):
                        os.remove(p)
                except Exception:
                    pass

    def _add_records(self, file_id: int, records: List[ParseResult]) -> int:
        n = 0
        with self.conn:
            for r in records:
                cur = self.conn.execute(
                    "INSERT INTO records(file_id,title,content,rtype,page,row_no,fields,extra) "
                    "VALUES(?,?,?,?,?,?,?,?)",
                    (file_id, r.title, r.content, r.rtype, r.page, r.row,
                     json.dumps(r.fields, ensure_ascii=False),
                     json.dumps(r.extra, ensure_ascii=False)))
                rid = cur.lastrowid
                self.conn.execute("INSERT INTO fts(rowid,title,content) VALUES(?,?,?)",
                                  (rid, fts_index_text(r.title), fts_index_text(r.content)))
                n += 1
            self.conn.execute("UPDATE files SET record_count=? WHERE id=?", (n, file_id))
        return n

    def import_file(self, src_path: str, rel_dir: str = "") -> Dict[str, Any]:
        """把 src_path 复制进来并解析索引。rel_dir 为可选目标子目录（相对 data/files）。

        同名文件重导规则（按 文件名 + 扩展名 + 所在文件夹 判定）：
        - 内容完全相同 → 跳过（不重复导入）；
        - 内容不同 → 替换旧版（删除旧记录与旧文件，避免出现两份）。
        """
        name = os.path.basename(src_path)
        ext = os.path.splitext(name)[1].lower()
        with open(src_path, "rb") as f:
            data = f.read()
        sha = hashlib.sha1(data).hexdigest()
        mtime = os.path.getmtime(src_path)

        src_abs = os.path.abspath(src_path)
        files_root = os.path.normcase(os.path.abspath(self.files_dir)) + os.sep
        in_files_dir = os.path.normcase(src_abs).startswith(files_root)

        # 计算目标相对路径（尚未写盘）
        if in_files_dir:
            # 文件本身就在 data/files 里：直接原位解析，避免重复复制
            rel = os.path.relpath(src_abs, self.root)
            dest = src_abs
        else:
            folder = _safe_rel_dir(rel_dir)
            stem = os.path.splitext(name)[0]
            base_name = "%s_%d%s" % (stem, int(time.time() * 1000), ext)
            rel = os.path.join(FILES_DIR, folder, base_name) if folder else os.path.join(FILES_DIR, base_name)
            dest = os.path.join(self.root, rel)
            base, ext2 = os.path.splitext(rel)
            i = 1
            while os.path.exists(dest):
                rel = "%s_%d%s" % (base, i, ext2)
                dest = os.path.join(self.root, rel)
                i += 1

        target_folder = os.path.normcase(os.path.normpath(os.path.dirname(rel)))

        # 同名检测（name + ext + 所在文件夹）
        dup_row = None
        replace_rows = []
        for r in self.conn.execute("SELECT id, path, sha FROM files WHERE name=? AND ext=?",
                                   (name, ext)).fetchall():
            if os.path.normcase(os.path.normpath(os.path.dirname(r["path"]))) != target_folder:
                continue
            if r["sha"] == sha:
                dup_row = r
            else:
                replace_rows.append(r)

        if dup_row is not None:
            return {"ok": True, "skipped": True, "file_id": dup_row["id"], "name": name, "records": 0}

        replaced = len(replace_rows)
        if replace_rows:
            # 替换旧版：删旧记录 + 旧 files 行 + 旧磁盘文件（但保留正在原位扫描的源文件）
            for r in replace_rows:
                with self.conn:
                    for o in self.conn.execute("SELECT id FROM records WHERE file_id=?", (r["id"],)):
                        self.conn.execute("DELETE FROM fts WHERE rowid=?", (o["id"],))
                    self.conn.execute("DELETE FROM records WHERE file_id=?", (r["id"],))
                    self.conn.execute("DELETE FROM files WHERE id=?", (r["id"],))
                old_path = os.path.join(self.root, r["path"])
                if os.path.normcase(os.path.abspath(old_path)) != os.path.normcase(src_abs) and os.path.isfile(old_path):
                    try:
                        os.remove(old_path)
                    except Exception:
                        pass

        # 写盘（web 上传才需要；原位扫描时 dest 即 src，已存在）
        if not in_files_dir:
            os.makedirs(os.path.dirname(dest), exist_ok=True)
            with open(dest, "wb") as f:
                f.write(data)

        file_id = self._insert_file(rel, name, ext, len(data), mtime, sha)
        res = self.reparse_file(file_id)
        if replaced:
            res["replaced"] = True
            res["replaced_files"] = replaced
        return res

    def reparse_file(self, file_id: int) -> Dict[str, Any]:
        row = self.conn.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()
        if not row:
            return {"ok": False, "error": "文件不存在"}
        with self.conn:
            old = self.conn.execute("SELECT id FROM records WHERE file_id=?", (file_id,)).fetchall()
            for o in old:
                self.conn.execute("DELETE FROM fts WHERE rowid=?", (o["id"],))
            self.conn.execute("DELETE FROM records WHERE file_id=?", (file_id,))
        path = os.path.join(self.root, row["path"])
        try:
            ext = row["ext"].lower()
            if ext in IMAGE_EXTS:
                records = parse_image_file(path)
            elif ext in EXT_PARSERS:
                records = EXT_PARSERS[ext](path)
            else:
                records = [ParseResult(row["name"], "", "other")]
            n = self._add_records(file_id, records)
            with self.conn:
                self.conn.execute("UPDATE files SET status='ok', error='' WHERE id=?", (file_id,))
            return {"ok": True, "file_id": file_id, "name": row["name"], "records": n}
        except Exception as e:
            with self.conn:
                self.conn.execute("UPDATE files SET status='error', error=? WHERE id=?", (str(e), file_id))
            return {"ok": False, "file_id": file_id, "name": row["name"], "error": str(e)}

    def scan_folder(self) -> Dict[str, Any]:
        """递归扫描 data/files（含子文件夹，如 标书素材/标书模板）。"""
        added = []
        replaced = []
        errors = []
        for dirpath, dirnames, filenames in os.walk(self.files_dir):
            dirnames.sort()
            for fname in sorted(filenames):
                p = os.path.join(dirpath, fname)
                if not os.path.isfile(p):
                    continue
                res = self.import_file(p)
                if res.get("ok") and not res.get("skipped"):
                    item = {"name": res["name"], "records": res.get("records", 0),
                            "folder": os.path.relpath(dirpath, self.files_dir)}
                    if res.get("replaced"):
                        replaced.append(item)
                    else:
                        added.append(item)
                elif not res.get("ok"):
                    errors.append({"name": fname, "error": res.get("error", "未知错误")})
        return {"added": added, "replaced": replaced, "errors": errors}

    def list_files(self) -> List[Dict[str, Any]]:
        rows = self.conn.execute(
            "SELECT f.*, (SELECT COUNT(*) FROM records r WHERE r.file_id=f.id) AS n "
            "FROM files f ORDER BY f.id DESC").fetchall()
        return [dict(r) for r in rows]

    def delete_file(self, file_id: int) -> Dict[str, Any]:
        self._delete_file(file_id)
        return {"ok": True}

    def stats(self) -> Dict[str, Any]:
        nf = self.conn.execute("SELECT COUNT(*) c FROM files").fetchone()["c"]
        nr = self.conn.execute("SELECT COUNT(*) c FROM records").fetchone()["c"]
        kinds = {}
        for r in self.conn.execute("SELECT kind, COUNT(*) c FROM files GROUP BY kind"):
            kinds[r["kind"]] = r["c"]
        return {"files": nf, "records": nr, "kinds": kinds, "db": self.db_path}
    def _fts_query(self, q: str) -> str:
        toks = cjk_tokens(q)
        return " ".join(toks) if toks else ""

    def search(self, q: str, limit: int = 30,
              kind: Optional[str] = None) -> List[Dict[str, Any]]:
        q = (q or "").strip()
        if not q:
            return []
        fts_q = self._fts_query(q)
        nq = norm_text(q)
        # 类型筛选：kind -> 允许的 file_kind 列表
        kinds = KIND_FILTERS.get(kind or "", [])
        kind_clause = ""
        kind_params: List[str] = []
        if kinds:
            kind_clause = " AND f.kind IN (%s)" % ",".join("?" * len(kinds))
            kind_params = kinds
        results: List[Dict[str, Any]] = []
        if fts_q and len(fts_q) >= 2:
            try:
                sql = ("SELECT r.id, r.title, r.content, r.rtype, r.page, r.row_no, r.fields, r.extra, "
                       "f.id AS file_id, f.name AS file_name, f.kind AS file_kind, "
                       "bm25(fts) AS bm FROM fts JOIN records r ON r.id=fts.rowid "
                       "JOIN files f ON f.id=r.file_id WHERE fts MATCH ?%s ORDER BY bm LIMIT ?" % kind_clause)
                rows = self.conn.execute(sql, (fts_q, *kind_params, max(limit * 5, 100))).fetchall()
                results = [dict(r) for r in rows]
            except sqlite3.OperationalError:
                results = []
        if len(results) < limit or not fts_q:
            like = "%" + q + "%"
            rows = self.conn.execute(
                "SELECT r.id, r.title, r.content, r.rtype, r.page, r.row_no, r.fields, r.extra, "
                "f.id AS file_id, f.name AS file_name, f.kind AS file_kind, 0 AS bm "
                "FROM records r JOIN files f ON f.id=r.file_id "
                "WHERE (r.title LIKE ? OR r.content LIKE ?)%s LIMIT ?" % kind_clause,
                (like, like, *kind_params, max(limit * 3, 60))).fetchall()
            seen = {r["id"] for r in results}
            for r in rows:
                if r["id"] not in seen:
                    results.append(dict(r))
                    seen.add(r["id"])
        scored = []
        for r in results:
            scored.append((self._score_result(r, q, nq), r))
        scored.sort(key=lambda x: -x[0])
        out = []
        for score, r in scored[:limit]:
            try:
                fields = json.loads(r["fields"] or "{}")
            except Exception:
                fields = {}
            try:
                extra = json.loads(r["extra"] or "{}")
            except Exception:
                extra = {}
            out.append({
                "id": r["id"], "title": r["title"], "content": r["content"],
                "type": r["rtype"], "page": r["page"], "row": r["row_no"],
                "fields": fields, "extra": extra,
                "file_id": r["file_id"], "file_name": r["file_name"], "file_kind": r["file_kind"],
                "score": round(score, 1),
                "snippet": highlight(r["content"] or r["title"], q),
            })
        return out

    def _score_result(self, r: Any, q: str, nq: str) -> float:
        score = 0.0
        nt = norm_text(r["title"])
        nc = norm_text(r["content"])
        if nq:
            if nq == nt:
                score += 100
            elif nq in nt:
                score += 80
            elif nt in nq:
                score += 60
            elif nq in nc:
                score += 40
        qt = set(cjk_tokens(q))
        tt = set(cjk_tokens(r["title"]))
        inter = len(qt & tt)
        union = len(qt | tt)
        if union:
            score += 25 * (inter / union)
        score += 12 * difflib.SequenceMatcher(None, nq, nt).ratio()
        bm = float(r["bm"] or 0)
        score += 10 * min(bm / 8.0, 1.0)
        if nq and nq in nc:
            score += 8
        return min(score, 130)

    def _field_keys(self) -> List[str]:
        keys: set = set()
        for r in self.conn.execute("SELECT fields FROM records WHERE fields IS NOT NULL AND fields != '{}'"):
            try:
                keys.update(json.loads(r["fields"]).keys())
            except Exception:
                pass
        return sorted(keys)

    def suggest_mapping(self, template_cols: List[str], title_col: str) -> Dict[str, str]:
        """为模板列建议 KB 字段映射（基于名称相似度）。"""
        kb_fields = self._field_keys()
        mapping: Dict[str, str] = {}
        for col in template_cols:
            if col == title_col:
                continue
            ncol = norm_text(col)
            best, best_score = None, 0.0
            for f in kb_fields:
                nf = norm_text(f)
                s = max(difflib.SequenceMatcher(None, ncol, nf).ratio(),
                        (1.0 if nf and (nf in ncol or ncol in nf) else 0.0))
                if s > best_score:
                    best, best_score = f, s
            if best and best_score >= 0.5:
                mapping[col] = best
        return mapping

    def match_title(self, q: str, top_k: int = 3) -> List[Dict[str, Any]]:
        """为单个标题词返回 top_k 个候选记录。"""
        q = (q or "").strip()
        if not q:
            return []
        candidates = self.search(q, limit=40)
        nq = norm_text(q)
        rows = self.conn.execute(
            "SELECT r.id, r.title, r.content, r.fields, f.name AS file_name "
            "FROM records r JOIN files f ON f.id=r.file_id WHERE r.title != '' LIMIT 2000").fetchall()
        seen = {c["id"] for c in candidates}
        for r in rows:
            if r["id"] in seen:
                continue
            nt = norm_text(r["title"])
            if nq and (nt == nq or nq in nt or nt in nq):
                candidates.append({"id": r["id"], "title": r["title"], "content": r["content"],
                                   "fields": {}, "file_name": r["file_name"], "score": 0, "bm": 0})
                seen.add(r["id"])
        scored = []
        for c in candidates:
            s = self._title_score(q, nq, c["title"])
            # 字段值匹配（表格的行标题可能不在第一列，字段值更贴近真实名称）
            try:
                fvals = json.loads(c.get("fields") or "{}").values()
            except Exception:
                fvals = []
            for v in fvals:
                if v:
                    s = max(s, self._title_score(q, nq, str(v)))
            # 正文包含
            nc = norm_text(c.get("content") or "")
            if nq and nq in nc:
                s = max(s, 50.0)
            bm = float(c.get("bm") or 0)
            s += 8 * min(bm / 8.0, 1.0)
            scored.append((s, c))
        scored.sort(key=lambda x: -x[0])
        out = []
        for s, c in scored[:top_k]:
            try:
                fields = json.loads(c.get("fields") or "{}")
            except Exception:
                fields = {}
            out.append({
                "record_id": c["id"], "title": c["title"],
                "file_name": c.get("file_name", ""),
                "content": (c.get("content") or "")[:200],
                "score": round(min(s, 100), 1), "fields": fields,
            })
        return out

    def _title_score(self, q: str, nq: str, title: str) -> float:
        nt = norm_text(title)
        if not nq or not nt:
            return 0
        if nq == nt:
            return 100
        if nq in nt:
            return 85 + 10 * (len(nq) / len(nt))
        if nt in nq:
            return 70 + 10 * (len(nt) / len(nq))
        qt = set(cjk_tokens(q))
        tt = set(cjk_tokens(title))
        inter = len(qt & tt)
        union = len(qt | tt)
        jac = inter / union if union else 0
        seq = difflib.SequenceMatcher(None, nq, nt).ratio()
        return min(80, 55 * jac + 35 * seq)

    def get_record(self, record_id: int) -> Optional[Dict[str, Any]]:
        r = self.conn.execute(
            "SELECT r.*, f.name AS file_name, f.kind AS file_kind FROM records r "
            "JOIN files f ON f.id=r.file_id WHERE r.id=?", (record_id,)).fetchone()
        if not r:
            return None
        d = dict(r)
        try:
            d["fields"] = json.loads(d["fields"] or "{}")
        except Exception:
            d["fields"] = {}
        return d

    def field_keys(self) -> List[str]:
        """所有表格字段名（只统计 kind=sheet 的记录，供表格工作台字段下拉使用）。"""
        keys: set = set()
        for r in self.conn.execute(
                "SELECT r.fields FROM records r JOIN files f ON f.id=r.file_id "
                "WHERE f.kind='sheet' AND r.fields IS NOT NULL AND r.fields!='{}'"):
            try:
                keys.update(json.loads(r["fields"]).keys())
            except Exception:
                pass
        return sorted(keys)

    def list_tables(self) -> List[Dict[str, Any]]:
        """表格工作台：列出所有表格文件（kind=sheet）及列名、行数。"""
        out: List[Dict[str, Any]] = []
        for f in self.conn.execute(
                "SELECT id, name, record_count FROM files "
                "WHERE kind='sheet' AND status='ok' ORDER BY name"):
            cols: List[str] = []
            r = self.conn.execute(
                "SELECT fields FROM records WHERE file_id=? AND fields IS NOT NULL AND fields!='{}' "
                "ORDER BY id LIMIT 1", (f["id"],)).fetchone()
            if r:
                try:
                    cols = list(json.loads(r["fields"]).keys())
                except Exception:
                    cols = []
            out.append({"file_id": f["id"], "name": f["name"],
                        "records": f["record_count"] or 0, "columns": cols})
        return out

    def sheet_rows(self, file_id: int, sort_field: str = "", sort_dir: str = "asc",
                   filter_field: str = "", filter_value: str = "",
                   limit: int = 200, offset: int = 0) -> Dict[str, Any]:
        """表格工作台：取某表格的行，支持按字段排序 / 过滤 / 分页。"""
        rows = self.conn.execute(
            "SELECT id, title, content, row_no, fields FROM records "
            "WHERE file_id=? AND rtype='sheet' ORDER BY id", (file_id,)).fetchall()
        fv_norm = norm_text(filter_value) if filter_value else ""
        items: List[Dict[str, Any]] = []
        for r in rows:
            try:
                fields = json.loads(r["fields"] or "{}")
            except Exception:
                fields = {}
            if filter_field:
                val = str(fields.get(filter_field, "") or "")
                if filter_value and not (filter_value in val or (fv_norm and fv_norm in norm_text(val))):
                    continue
            items.append({"record_id": r["id"], "title": r["title"], "row": r["row_no"],
                          "fields": fields, "content": (r["content"] or "")[:800]})
        if sort_field:
            rev = (sort_dir or "").lower() == "desc"

            def _key(x: Dict[str, Any]):
                v = x["fields"].get(sort_field, "")
                s = str(v if v is not None else "")
                try:
                    return (0, float(s))
                except Exception:
                    return (1, norm_text(s))

            items.sort(key=_key, reverse=rev)
        total = len(items)
        return {"total": total, "rows": items[offset:offset + limit]}

    def search_by_field(self, field: str, value: str, limit: int = 50) -> List[Dict[str, Any]]:
        """表格工作台：跨所有表格按字段值筛选（字段值包含 value）。"""
        field = (field or "").strip()
        value = (value or "").strip()
        if not field or not value:
            return []
        v_norm = norm_text(value)
        out: List[Dict[str, Any]] = []
        for r in self.conn.execute(
                "SELECT r.id, r.title, r.fields, f.id AS file_id, f.name AS file_name "
                "FROM records r JOIN files f ON f.id=r.file_id "
                "WHERE r.fields IS NOT NULL AND r.fields!='{}' AND f.kind='sheet' ORDER BY r.id"):
            try:
                fields = json.loads(r["fields"] or "{}")
            except Exception:
                continue
            if field not in fields:
                continue
            fv = str(fields[field] if fields[field] is not None else "")
            if value in fv or (v_norm and v_norm in norm_text(fv)):
                out.append({"record_id": r["id"], "title": r["title"], "file_id": r["file_id"],
                            "file_name": r["file_name"], "fields": fields, "match_value": fv})
                if len(out) >= limit:
                    break
        return out

    def record_pages(self, record_id: int) -> List[Dict[str, Any]]:
        """阅读模式：返回同一文件下所有带页码的记录（用于 PDF/PPT 翻页）。"""
        r = self.conn.execute("SELECT file_id FROM records WHERE id=?", (record_id,)).fetchone()
        if not r:
            return []
        rows = self.conn.execute(
            "SELECT id, title, page FROM records WHERE file_id=? AND page IS NOT NULL "
            "ORDER BY page", (r["file_id"],)).fetchall()
        return [{"id": x["id"], "title": x["title"], "page": x["page"]} for x in rows]

    def get_file_path(self, file_id: int) -> Optional[str]:
        """返回文件的绝对路径（用于图片预览等）。"""
        r = self.conn.execute("SELECT path FROM files WHERE id=?", (file_id,)).fetchone()
        if not r:
            return None
        return os.path.join(self.root, r["path"])

    def export_selected(self, record_ids: List[int],
                         scores: Optional[Dict[int, float]] = None) -> str:
        """把选中的记录导出为 xlsx。"""
        if openpyxl is None:
            raise RuntimeError("缺少 openpyxl 库")
        out_path = os.path.join(self.export_dir,
                                "搜索导出_%s.xlsx" % _dt.datetime.now().strftime("%Y%m%d_%H%M%S"))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "搜索结果"
        ws.append(["匹配度", "标题", "内容", "来源文件", "类型", "页码", "行号", "字段(JSON)"])
        for rid in record_ids:
            rec = self.get_record(rid)
            if not rec:
                continue
            sc = (scores or {}).get(rid, "")
            ws.append([sc, rec.get("title", ""), rec.get("content", ""),
                       rec.get("file_name", ""), rec.get("rtype", ""), rec.get("page", ""),
                       rec.get("row_no", ""), json.dumps(rec.get("fields", {}), ensure_ascii=False)])
        _autofit(ws)
        wb.save(out_path)
        return out_path

    def export_filled(self, template_path: str, choices: Dict[int, Optional[int]],
                      mapping: Dict[str, str], title_col: str,
                      scores: Optional[Dict[int, float]] = None) -> str:
        """生成填写结果 xlsx。choices: {模板行号(0-based): record_id 或 None}。"""
        if pd is None or openpyxl is None:
            raise RuntimeError("缺少 pandas/openpyxl 库")
        ext = os.path.splitext(template_path)[1].lower()
        if ext == ".csv":
            try:
                df = pd.read_csv(template_path, encoding="utf-8-sig")
            except UnicodeDecodeError:
                df = pd.read_csv(template_path, encoding="gb18030")
        else:
            df = pd.read_excel(template_path, sheet_name=0)
        df = df.where(df.notna(), "")
        out = df.copy()
        matched_col = "匹配结果"
        source_col = "匹配来源"
        score_col = "匹配度"
        for c in mapping:
            if c in out.columns:
                out[c] = ""
        if matched_col not in out.columns:
            out[matched_col] = ""
        if source_col not in out.columns:
            out[source_col] = ""
        if score_col not in out.columns:
            out[score_col] = ""
        # 避免 pandas 严格类型导致写入混合类型失败
        out[matched_col] = out[matched_col].astype(object)
        out[source_col] = out[source_col].astype(object)
        out[score_col] = out[score_col].astype(object)
        for c in mapping:
            if c in out.columns:
                out[c] = out[c].astype(object)
        for idx in df.index:
            rec_id = choices.get(int(idx))
            if rec_id is None:
                out.at[idx, matched_col] = "未匹配"
                continue
            rec = self.get_record(rec_id)
            if not rec:
                out.at[idx, matched_col] = "未匹配"
                continue
            fields = rec.get("fields", {})
            for c, f in mapping.items():
                if c in out.columns:
                    out.at[idx, c] = "" if fields.get(f) is None else fields.get(f)
            out.at[idx, matched_col] = "已匹配"
            out.at[idx, source_col] = rec.get("file_name", "")
            out.at[idx, score_col] = (scores or {}).get(int(idx), "")
        out_path = os.path.join(self.export_dir,
                                "填写结果_%s.xlsx" % _dt.datetime.now().strftime("%Y%m%d_%H%M%S"))
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            out.to_excel(writer, index=False, sheet_name="填写结果")
        wb = openpyxl.load_workbook(out_path)
        _autofit(wb.active)
        wb.save(out_path)
        return out_path

    def close(self):
        try:
            self.conn.close()
        except Exception:
            pass

    def paste_preview(self, source_path: str, target_path: str, src_key: str,
                      tgt_key: str, top_k: int = 3, mode: str = "match") -> Dict[str, Any]:
        """表格粘贴预览：读取源表/目标表，给出映射建议与（按匹配列的）逐行候选。"""
        src = read_template(source_path)
        tgt = read_template(target_path)
        if not src.get("ok"):
            return {"ok": False, "error": "源表格读取失败：" + src.get("error", "")}
        if not tgt.get("ok"):
            return {"ok": False, "error": "目标表格读取失败：" + tgt.get("error", "")}
        scols, tcols = src["columns"], tgt["columns"]
        if not src_key or src_key not in scols:
            src_key = scols[0]
        if not tgt_key or tgt_key not in tcols:
            tgt_key = tcols[0]
        mapping: Dict[str, str] = {}
        for tc in tcols:
            if tc == tgt_key:
                continue
            ntc = norm_text(tc)
            best, bs = None, 0.0
            for sc in scols:
                nsc = norm_text(sc)
                s = max(difflib.SequenceMatcher(None, ntc, nsc).ratio(),
                        (1.0 if nsc and (nsc in ntc or ntc in nsc) else 0.0))
                if s > bs:
                    best, bs = sc, s
            if best and bs >= 0.5:
                mapping[tc] = best
        rows_out: List[Dict[str, Any]] = []
        if mode == "match":
            for i, row in enumerate(tgt["rows"]):
                q = _clean_cell(row.get(tgt_key, ""))
                rows_out.append({"index": i, "key": q,
                                 "candidates": self._top_source_rows(src, src_key, q, top_k)})
        return {"ok": True, "mode": mode,
                "source_columns": scols, "target_columns": tcols,
                "src_key": src_key, "tgt_key": tgt_key,
                "mapping": mapping, "rows": rows_out,
                "source_rows": len(src["rows"]), "target_rows": len(tgt["rows"])}


    def _top_source_rows(self, src: Dict[str, Any], src_key: str, q: str,
                         top_k: int) -> List[Dict[str, Any]]:
        nq = norm_text(q)
        scored = []
        for i, row in enumerate(src["rows"]):
            v = _clean_cell(row.get(src_key, ""))
            s = self._title_score(q, nq, v) if v else 0.0
            scored.append((s, i))
        scored.sort(key=lambda x: -x[0])
        out = []
        for s, i in scored[:top_k]:
            row = src["rows"][i]
            out.append({"source_index": i, "key": _clean_cell(row.get(src_key, "")),
                        "score": round(min(s, 100), 1),
                        "fields": {k: _clean_cell(v) for k, v in row.items()}})
        return out


    def paste_apply(self, source_path: str, target_path: str, mode: str,
                    choices: Dict[int, Optional[int]], mapping: Dict[str, str],
                    src_key: str, tgt_key: str, overwrite: bool = False,
                    append_unmatched: bool = True) -> str:
        """生成粘贴结果 xlsx。
        mode="match"：choices = {目标行号: 源行号 或 None}
        mode="row"  ：按行顺序直接复制，choices 忽略。"""
        if pd is None or openpyxl is None:
            raise RuntimeError("缺少 pandas/openpyxl 库")
        src = read_template(source_path)
        tgt = read_template(target_path)
        if not src.get("ok") or not tgt.get("ok"):
            raise RuntimeError("表格读取失败")
        scols, tcols = src["columns"], tgt["columns"]
        if not src_key or src_key not in scols:
            src_key = scols[0]
        if not tgt_key or tgt_key not in tcols:
            tgt_key = tcols[0]
        out_cols = list(tcols)
        for tc in mapping:
            if tc not in out_cols:
                out_cols.append(tc)
        valid_map = {tc: sc for tc, sc in mapping.items()
                     if tc in out_cols and sc in scols and tc != tgt_key}
        out_rows: List[Dict[str, Any]] = []
        chosen_src: set = set()
        if mode == "row":
            n = min(len(src["rows"]), len(tgt["rows"]))
            for i, row in enumerate(tgt["rows"]):
                r = {c: _clean_cell(row.get(c, "")) for c in tcols}
                if i < n:
                    for tc, sc in valid_map.items():
                        if overwrite or not _clean_cell(r.get(tc, "")):
                            r[tc] = _clean_cell(src["rows"][i].get(sc, ""))
                out_rows.append(r)
        else:
            for i, row in enumerate(tgt["rows"]):
                r = {c: _clean_cell(row.get(c, "")) for c in tcols}
                src_idx = choices.get(i)
                if src_idx is not None and 0 <= src_idx < len(src["rows"]):
                    srow = src["rows"][src_idx]
                    for tc, sc in valid_map.items():
                        if overwrite or not _clean_cell(r.get(tc, "")):
                            r[tc] = _clean_cell(srow.get(sc, ""))
                    chosen_src.add(src_idx)
                out_rows.append(r)
        if append_unmatched:
            for i, srow in enumerate(src["rows"]):
                if mode == "row" and i < len(tgt["rows"]):
                    continue
                if mode == "match" and i in chosen_src:
                    continue
                r = {c: "" for c in tcols}
                if tgt_key in out_cols:
                    r[tgt_key] = _clean_cell(srow.get(src_key, ""))
                for tc, sc in valid_map.items():
                    r[tc] = _clean_cell(srow.get(sc, ""))
                out_rows.append(r)
        out_path = os.path.join(self.export_dir,
                                "粘贴结果_%s.xlsx" % _dt.datetime.now().strftime("%Y%m%d_%H%M%S"))
        df = pd.DataFrame(out_rows, columns=out_cols)
        df = df.where(df.notna(), "")
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="粘贴结果")
        wb = openpyxl.load_workbook(out_path)
        _autofit(wb.active)
        wb.save(out_path)
        return out_path




    def bid_dir(self) -> str:
        d = os.path.join(self.root, "data", "bid")
        os.makedirs(d, exist_ok=True)
        return d

    def get_company_profile(self) -> Dict[str, Any]:
        path = os.path.join(self.bid_dir(), "company.json")
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict) and "fields" in data:
                return data
        except Exception:
            pass
        return {"fields": {}, "updated": ""}

    def save_company_profile(self, fields: Dict[str, Any]) -> Dict[str, Any]:
        path = os.path.join(self.bid_dir(), "company.json")
        clean = {k: _clean_cell(v) for k, v in fields.items() if k}
        data = {"fields": clean, "updated": now_str()}
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return data

    def _match_chapter_materials(self, title: str, top_k: int = 3) -> List[Dict[str, Any]]:
        """按章节标题（已去序号）在资料里模糊匹配素材：标题/内容包含 + 关键词联想。"""
        keywords = {
            "业绩": ["业绩", "项目", "合同"], "资质": ["资质", "证书", "资格"],
            "资格": ["资质", "证书", "资格"], "报价": ["报价", "价格", "金额"],
            "商务": ["报价", "价格", "金额", "商务"], "技术": ["技术", "方案", "服务", "系统"],
            "方案": ["方案", "技术", "服务", "系统"], "承诺": ["承诺", "售后", "服务"],
            "售后": ["承诺", "售后", "服务"], "简介": ["公司", "简介"],
            "公司": ["公司", "简介"], "营业执照": ["营业执照", "执照"],
            "法定代表人": ["法定代表人", "法人"], "身份证明": ["法定代表人", "身份", "法人"],
        }
        terms = set()
        nt = norm_text(title)
        if len(nt) >= 2:
            terms.add(nt)
        for k, exts in keywords.items():
            if k in nt:
                terms.update(exts)
        terms = [t for t in terms if len(t) >= 2]
        if not terms:
            return []
        rows = self.conn.execute(
            "SELECT r.id, r.title, r.content, f.name AS file_name FROM records r "
            "JOIN files f ON f.id=r.file_id").fetchall()
        scored = []
        for r in rows:
            ntitle = norm_text(r["title"])
            ncontent = norm_text(r["content"])
            s = 0.0
            for t in terms:
                if ntitle == t:
                    s = max(s, 100.0)
                elif t in ntitle:
                    s = max(s, 92.0)
                elif ntitle and ntitle in t:
                    s = max(s, 85.0)
                if t in ncontent:
                    s = max(s, 62.0)
            # 序列相似度加分：让“技术服务方案”这类近似标题排到前面
            if nt and len(nt) >= 2:
                seq = difflib.SequenceMatcher(None, nt, ntitle).ratio()
                s += 15 * seq
            if s > 0:
                scored.append((s, r))
        scored.sort(key=lambda x: -x[0])
        out = []
        for s, r in scored[:top_k]:
            out.append({"record_id": r["id"], "title": r["title"],
                        "file_name": r["file_name"], "score": round(s, 1)})
        return out


    def parse_bid_template(self, path: str) -> Dict[str, Any]:
        """解析 docx 标书模板：按标题拆出章节，并为每章在资料里找候选素材。"""
        if _docx_mod is None:
            return {"ok": False, "error": "缺少 python-docx 库"}
        d = _docx_mod.Document(path)
        chapters: List[Dict[str, Any]] = []
        cur = {"index": 0, "level": 1, "title": "", "template_text": ""}

        def flush():
            if cur["title"]:
                chapters.append(dict(cur))

        for para in d.paragraphs:
            style = (para.style.name or "").lower()
            text = para.text.strip()
            ol = getattr(para.paragraph_format, "outline_level", None)
            is_heading = (style.startswith("heading")
                          or "标题" in style
                          or (ol is not None and ol < 9))
            if is_heading and text:
                flush()
                m = re.match(r"heading\s*(\d)", style)
                level = int(m.group(1)) if m else (2 if "标题 2" in style else 1)
                cur = {"index": len(chapters), "level": min(level, 3),
                       "title": text, "template_text": ""}
            elif text:
                cur["template_text"] += text + "\n"
        flush()
        tpl_name = os.path.basename(path)
        out = []
        for ch in chapters:
            clean = re.sub(r"^(第[一二三四五六七八九十百0-9]+章\s*|\d+(\.\d+)*\s*)", "", ch["title"]).strip()
            cands = self._match_chapter_materials(clean, 3) if clean else []
            cands = [c for c in cands if c["file_name"] != tpl_name]
            out.append({**ch, "candidates": cands})
        return {"ok": True, "chapters": out, "total": len(out)}

    def _fill_tokens(self, text: str, profile: Dict[str, str]) -> str:
        """把文本里的 【字段】 / {{字段}} 按公司档案替换。"""
        def rep(m):
            tok = (m.group(1) or m.group(2) or "").strip()
            ntok = norm_text(tok)
            best, bs = None, 0.0
            for k, v in profile.items():
                nk = norm_text(k)
                s = max(difflib.SequenceMatcher(None, ntok, nk).ratio(),
                        (1.0 if nk and (nk in ntok or ntok in nk) else 0.0))
                if s > bs:
                    best, bs = (_clean_cell(v), s)
            return best if bs >= 0.6 else m.group(0)
        return re.sub(r"【([^】]+)】|\{\{([^}]+)\}\}", rep, text)

    def generate_bid(self, template_path: str, choices: Dict[str, Any],
                     project: Dict[str, Any]) -> str:
        """按模板章节生成投标文件 docx 草稿。
        choices: {章节index: {mode: auto/template/empty, record_id: 可选}}"""
        if _docx_mod is None:
            raise RuntimeError("缺少 python-docx 库")
        parsed = self.parse_bid_template(template_path)
        if not parsed.get("ok"):
            raise RuntimeError(parsed.get("error", "模板解析失败"))
        profile = self.get_company_profile()["fields"]
        fill_map = dict(profile)
        fill_map.update({k: _clean_cell(v) for k, v in project.items() if v})
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = _docx_mod.Document()

        def center(text, size=18, bold=True):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)
            return p

        proj_name = _clean_cell(project.get("项目名称") or "")
        if proj_name:
            center(proj_name, 22)
        center("投  标  文  件", 26)
        doc.add_paragraph()
        center("（草稿 · 请人工核对后投标）", 12, bold=False)
        doc.add_paragraph()
        for k, key in (("招标编号", "招标编号"), ("招标人", "招标人"),
                       ("投标人", "投标人"), ("日期", "日期")):
            v = ""
            if k == "投标人":
                v = profile.get("公司名称", "")
            elif k == "日期":
                v = _dt.date.today().strftime("%Y年%m月%d日")
            else:
                v = _clean_cell(project.get(k) or "")
            if v:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run("%s：%s" % (k, v))
        doc.add_page_break()
        doc.add_heading("目  录", level=1)
        for ch in parsed["chapters"]:
            doc.add_paragraph(ch["title"])
        doc.add_page_break()
        for ch in parsed["chapters"]:
            c = choices.get(str(ch["index"])) or {}
            mode = c.get("mode", "auto")
            doc.add_heading(ch["title"], level=min(ch.get("level", 1), 3))
            if mode == "auto":
                rid = c.get("record_id")
                rec = self.get_record(rid) if rid else None
                if rec:
                    fields = rec.get("fields") or {}
                    for k, v in fields.items():
                        if _clean_cell(v):
                            doc.add_paragraph("%s：%s" % (k, self._fill_tokens(v, fill_map)))
                    for line in (rec.get("content") or "").splitlines():
                        if line.strip():
                            doc.add_paragraph(self._fill_tokens(line, fill_map))
                else:
                    doc.add_paragraph("（未选择素材，待人工补充）")
            elif mode == "template":
                for line in ch["template_text"].splitlines():
                    if line.strip():
                        doc.add_paragraph(self._fill_tokens(line, fill_map))
            else:
                doc.add_paragraph("（本章留空，待人工填写）")
        out_path = os.path.join(self.export_dir,
                                "标书草稿_%s.docx" % _dt.datetime.now().strftime("%Y%m%d_%H%M%S"))
        doc.save(out_path)
        return out_path

def _autofit(ws, max_w: int = 50):
    for col in ws.columns:
        vals = [c.value for c in col if c.value is not None]
        width = max(len(str(v)) for v in vals) if vals else 10
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), max_w)


def read_template(path: str) -> Dict[str, Any]:
    """读取模板表格：返回列名、每行各列的值。"""
    if pd is None:
        return {"ok": False, "error": "缺少 pandas 库"}
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".csv":
            df = pd.read_csv(path, encoding="utf-8-sig")
        else:
            df = pd.read_excel(path, sheet_name=0)
    except UnicodeDecodeError:
        df = pd.read_csv(path, encoding="gb18030")
    df = df.where(df.notna(), "")
    cols = [str(c) for c in df.columns]
    rows = []
    for idx, r in df.iterrows():
        rows.append({str(c): _clean_cell(r[c]) for c in df.columns})
    return {"ok": True, "columns": cols, "rows": rows}


KIND_FILTERS = {
    "sheet": ["sheet"],                # 表格（xlsx/xls/csv）
    "doc": ["doc", "text"],            # 文档（Word/文本）
    "pdf": ["pdf"],                    # PDF
    "slide": ["slide"],                # 演示（pptx/ppt）
    "image": ["image"],                # 图片
}


ENV_CHECKS = [
    {"import": "pandas",       "pip": "pandas",       "label": "表格解析 / 自动填写（xlsx/csv）", "module": pd},
    {"import": "openpyxl",     "pip": "openpyxl",     "label": "Excel 解析（xlsx）",              "module": openpyxl},
    {"import": "xlrd",         "pip": "xlrd",         "label": "旧版表格（.xls）",               "module": xlrd},
    {"import": "pdfplumber",   "pip": "pdfplumber",   "label": "PDF 解析",                       "module": pdfplumber},
    {"import": "docx",         "pip": "python-docx",  "label": "Word 解析（.docx）",             "module": _docx_mod},
    {"import": "pptx",         "pip": "python-pptx",  "label": "PowerPoint 解析（.pptx）",       "module": _pptx_mod},
    {"import": "PIL",          "pip": "Pillow",       "label": "图片信息（png/jpg 等）",         "module": Image},
    {"import": "olefile",      "pip": "olefile",      "label": "旧版 Word/PPT（.doc/.ppt）",     "module": olefile},
]

INSTALL_COMMAND = ("pip install pandas openpyxl xlrd pdfplumber python-docx "
                   "python-pptx Pillow olefile")


def check_environment() -> List[Dict[str, Any]]:
    """检查本机解析组件，返回缺失项列表。每项含 import/pip/label/installed。"""
    missing = []
    for c in ENV_CHECKS:
        item = dict(c)
        item["installed"] = c["module"] is not None
        if not item["installed"]:
            missing.append(item)
    return missing


def build_env_message(missing: List[Dict[str, Any]]) -> str:
    lines = ["检测到本机缺少以下解析组件，对应格式将无法导入：\n"]
    for c in missing:
        lines.append("  • %s（需要安装 %s）" % (c["label"], c["pip"]))
    lines.append("\n基础功能（txt / csv 搜索）不受影响。")
    lines.append("\n解决办法（任选其一）：")
    lines.append("  1) 双击程序目录里的 install_deps.bat（推荐）")
    lines.append("  2) 在命令行执行：")
    lines.append("     " + INSTALL_COMMAND)
    return "\n".join(lines)



    # ---------- 标书助手 ----------
def get_python_info() -> Dict[str, Any]:
    import sys
    return {
        "python": sys.version.split()[0],
        "pandas": bool(pd), "openpyxl": bool(openpyxl), "pdfplumber": bool(pdfplumber),
        "docx": bool(_docx_mod), "pptx": bool(_pptx_mod), "pillow": bool(Image),
        "pytesseract": bool(pytesseract),
        "xlrd": bool(xlrd), "olefile": bool(olefile),
    }
